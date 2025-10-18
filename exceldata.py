import openpyxl
import os
import logging


rowdata = None
row_num = None
logging.basicConfig(level = logging.INFO, format= "%(asctime)s-%(levelname)s-%(message)s")
logger = logging.getLogger(__name__)
print('-'*60)
print("Start the datadriven excel")
print('-'*60)


def excel_function(file_name, sheet_name):
    try:
        if not os.path.exists(file_name):
          logger.error(f"The {file_name} is not found")
          return None
        else:
          wb = openpyxl.load_workbook(file_name)
          logger.info(f"The {file_name} is loaded")
    except Exception as e:
        logger.error(f"The error is displayed while loading : {e}")
        return None
        

    try:
        if sheet_name not in wb.sheetnames:
          print(sheet_name)
          logger.error(f"The sheet name {sheet_name} not found")
          return None
        else:
           sheet = wb[sheet_name]
           logger.info(f"The sheet name {sheet_name} is loaded")
    except Exception as e:
       logger.error(f"The excel sheet {sheet_name} is not loaded : {e}")
       return None

    try:
        heading_dict = {}
        column_data = sheet.cell(row = 1, column = 1).value
        if column_data is None:
          logger.error("There is no heading data")
          return None
        else:
          for col in range(1, sheet.max_column+1):
             heading_name = sheet.cell(row = 1, column = col).value
             heading_dict[heading_name]= col
          logger.info("The heading name is added to heading dictionary")
    except Exception as e:
       logger.error(f"The error is displayed while heading name added to heading dictionary : {e}")
       return None
    
    for row_num in range(2, sheet.max_row+1):
        rowdata = sheet.cell(row =2, column =1).value
    try:
          if rowdata is None:
             logger.error("There is no row data")
          else:
             logger.info("There is data in the row")
    except Exception as e:
           logger.error(f"The error displayed while checking the row data:{e}")

    try:
        if "Status" not in heading_dict:
               logger.error("The Status name is not displayed on column heading")
        else:
               logger.info("The Status column is displayed on column heading")
    except Exception as e:
          logger.error(f"The error while checking status column:{e}")
          return None

    try:
        if sheet.cell(row = row_num, column = heading_dict["Status"]).value == "Used":
            logger.info(f"The {row_num} number row is alredy field skip the row")
    except Exception as e:
        logger.error(f"The erro displayed while checking alredy used :{e}")
        return None
    
 

          
        
        

    
          
        # if "Status" in heading_dict:
        #   continue
        # if sheet.cell(row = row, column = heading_dict["Status"]).value == "Used"
       
          
       
       
          
          

    print('-'*60)
    print("Complete the datadriven excel")
    print('-'*60)

excel_function("text_excel.xlsx", "Form")
